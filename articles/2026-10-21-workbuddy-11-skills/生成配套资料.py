"""从 技能清单.json 生成配套资料：Markdown 清单与 Excel 表格。依赖 openpyxl。"""
from pathlib import Path
import json

ROOT = Path(__file__).resolve().parent
OUT = ROOT / '配套资料'


def load():
    lib = json.loads((ROOT / '技能清单.json').read_text(encoding='utf-8'))
    assert [s['id'] for s in lib['skills']] == list(range(1, 12))
    return lib


def markdown_list(lib):
    OUT.mkdir(exist_ok=True)
    lines = [f"# {lib['title']}", '', f"版本 {lib['version']} · 更新 {lib['updated']}", '', lib['note'], '']
    for s in lib['skills']:
        step = f" · 流水线：{s['step']}" if s.get('step') else ''
        lines += [f"## {s['id']:02d}｜{s['name']}（{s['install']}）", '', f"分组：{s['group']}{step} · 来源：{s['source']}", '',
                  s['what'], '', '怎么说：', '', '```text', s['prompt'], '```', '', f"注意：{s['caution']}", '', s['url'], '']
    (OUT / 'WorkBuddy好用技能11个.md').write_text('\n'.join(lines), encoding='utf-8')


def excel(lib):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    OUT.mkdir(exist_ok=True)
    wb = Workbook(); ws = wb.active; ws.title = '技能清单'
    ws.append(['编号', '技能', '安装名', '分组', '流水线步骤', '来源', '能做什么', '怎么说', '注意', '地址'])
    for s in lib['skills']:
        ws.append([s['id'], s['name'], s['install'], s['group'], s.get('step', ''), s['source'], s['what'], s['prompt'], s['caution'], s['url']])
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='4457C4')
    for col, w in zip('ABCDEFGHIJ', [6, 18, 20, 8, 10, 10, 40, 60, 40, 50]):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical='top')
    ws.freeze_panes = 'A2'
    wb.save(OUT / 'WorkBuddy好用技能11个.xlsx')


if __name__ == '__main__':
    lib = load(); markdown_list(lib); excel(lib)
    print('已生成：WorkBuddy好用技能11个.md / .xlsx')
