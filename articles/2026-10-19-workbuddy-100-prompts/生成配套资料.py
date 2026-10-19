"""从 提示词库.json 生成配套资料：Markdown 清单、Excel 表格、可搜索网页。依赖 openpyxl。"""
from pathlib import Path
import json

ROOT = Path(__file__).resolve().parent
OUT = ROOT / '配套资料'
CN = '一二三四五六七八九十'


def load():
    lib = json.loads((ROOT / '提示词库.json').read_text(encoding='utf-8'))
    ids = [p['id'] for c in lib['categories'] for p in c['prompts']]
    assert ids == list(range(1, 101)), '编号必须是 1—100 连续'
    assert all(len(c['prompts']) == 10 for c in lib['categories'])
    return lib


def markdown_list(lib):
    lines = [f"# {lib['title']}", '', f"版本 {lib['version']} · 更新 {lib['updated']} · 作者 {lib['author']}", '',
             f"使用方法：{lib['usage']}", '']
    for cat in lib['categories']:
        lines += [f"## {CN[cat['id']-1]}、{cat['name']}", '', cat['intro'], '']
        for p in cat['prompts']:
            lines += [f"### {p['id']:02d}｜{p['title']}", '', '```text', p['prompt'], '```', '']
    (OUT / 'WorkBuddy提示词100例.md').write_text('\n'.join(lines), encoding='utf-8')


def excel(lib):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook(); ws = wb.active; ws.title = '提示词100例'
    ws.append(['编号', '分类', '用途', '提示词'])
    for cat in lib['categories']:
        for p in cat['prompts']:
            ws.append([p['id'], cat['name'], p['title'], p['prompt']])
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='4457C4')
    for col, w in zip('ABCD', [7, 18, 24, 100]):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical='top')
    ws.freeze_panes = 'A2'
    wb.save(OUT / 'WorkBuddy提示词100例.xlsx')


def web(lib):
    data = json.dumps(lib, ensure_ascii=False)
    page = (ROOT / '网页模板.html').read_text(encoding='utf-8').replace('/*DATA*/null', data)
    (OUT / 'WorkBuddy提示词100例.html').write_text(page, encoding='utf-8')


def main():
    lib = load()
    OUT.mkdir(exist_ok=True)
    markdown_list(lib); excel(lib); web(lib)
    print('已生成：WorkBuddy提示词100例.md / .xlsx / .html')


if __name__ == '__main__':
    main()
